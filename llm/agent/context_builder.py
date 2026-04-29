import io
from pathlib import Path

from docx import Document as DocxDocument
from PyPDF2 import PdfReader
import openpyxl
from pptx import Presentation
from django.db.models import Q

from main.models import PublicationFile, PublicationReference


class ContextBuilder:
    def __init__(self, project,  user, document_text_limit=7500):
        self.project = project
        self.user = user
        self.document_text_limit = document_text_limit

    # ----------------------------
    # Universal safe text cleaner
    # ----------------------------
    def _clean_text(self, value):
        if value is None:
            return ""

        text = str(value).replace("\x00", "").strip()
        return text

    # ----------------------------
    # DOCX parser
    # ----------------------------
    def parse_docx(self, file_path):
        try:
            doc = DocxDocument(file_path)
            chunks = []

            for paragraph in doc.paragraphs:
                text = self._clean_text(paragraph.text)
                if text:
                    chunks.append(text)

            for table in doc.tables:
                for row in table.rows:
                    row_values = []
                    for cell in row.cells:
                        cell_text = self._clean_text(cell.text)
                        if cell_text:
                            row_values.append(cell_text)

                    if row_values:
                        chunks.append(" | ".join(row_values))

            return "\n".join(chunks)

        except Exception as exc:
            return f"[DOCX parse error: {exc}]"

    # ----------------------------
    # PDF parser
    # ----------------------------
    def parse_pdf(self, file_path):
        try:
            reader = PdfReader(file_path)
            chunks = []

            for page in reader.pages:
                text = page.extract_text() or ""
                text = self._clean_text(text)
                if text:
                    chunks.append(text)

            return "\n".join(chunks)
        except Exception as exc:
            return f"[PDF parse error: {exc}]"

    # ----------------------------
    # TXT parser
    # ----------------------------
    def parse_txt(self, file_path):
        encodings = ["utf-8", "utf-8-sig", "cp1251", "latin-1"]

        for enc in encodings:
            try:
                with open(file_path, "r", encoding=enc) as f:
                    return self._clean_text(f.read())
            except Exception:
                continue

        return "[TXT parse error: unable to decode file]"

    # ----------------------------
    # Excel parser
    # ----------------------------
    def parse_excel(self, file_path):
        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            chunks = []

            for sheet in wb.worksheets:
                chunks.append(f"# Sheet: {sheet.title}")

                for row in sheet.iter_rows(values_only=True):
                    values = []
                    for cell in row:
                        if cell is not None:
                            values.append(self._clean_text(cell))

                    if values:
                        chunks.append(" | ".join(values))

            return "\n".join(chunks)

        except Exception as exc:
            return f"[Excel parse error: {exc}]"

    # ----------------------------
    # PowerPoint parser
    # ----------------------------
    def parse_pptx(self, file_path):
        try:
            prs = Presentation(file_path)
            chunks = []

            for index, slide in enumerate(prs.slides, start=1):
                chunks.append(f"# Slide {index}")

                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text = self._clean_text(shape.text)
                        if text:
                            chunks.append(text)

            return "\n".join(chunks)

        except Exception as exc:
            return f"[PPTX parse error: {exc}]"

    # ----------------------------
    # Fallback parser
    # ----------------------------
    def parse_other(self, file_path):
        try:
            suffix = Path(file_path).suffix.lower()

            if suffix in [".json", ".csv", ".md", ".log", ".xml", ".html"]:
                return self.parse_txt(file_path)

            return "[Unsupported file type]"

        except Exception as exc:
            return f"[Other parse error: {exc}]"

    # ----------------------------
    # Dispatcher
    # ----------------------------
    def _resolve_document_type(self, document):
        explicit_type = str(getattr(document, "file_type", "") or "").strip().lower()
        if explicit_type in {"docx", "pdf", "txt", "excel", "pptx"}:
            return explicit_type

        file_name = ""
        if getattr(document, "file", None):
            file_name = getattr(document.file, "name", "") or getattr(document.file, "path", "") or ""
        suffix = Path(file_name).suffix.lower()
        extension_map = {
            ".docx": "docx",
            ".pdf": "pdf",
            ".txt": "txt",
            ".xlsx": "excel",
            ".xls": "excel",
            ".csv": "excel",
            ".pptx": "pptx",
            ".ppt": "pptx",
        }
        return extension_map.get(suffix, "other")

    def parse_document_content(self, document, text_limit=None):
        limit = text_limit if text_limit is not None else self.document_text_limit
        if not document:
            return ""

        if not document.file:
            return self._clean_text(getattr(document, "content", ""))

        file_path = document.file.path
        file_type = self._resolve_document_type(document)

        if file_type == "docx":
            return self.parse_docx(file_path)[:limit]

        if file_type == "pdf":
            return self.parse_pdf(file_path)[:limit]

        if file_type == "txt":
            return self.parse_txt(file_path)[:limit]

        if file_type == "excel":
            return self.parse_excel(file_path)[:limit]

        if file_type == "pptx":
            return self.parse_pptx(file_path)[:limit]

        return self.parse_other(file_path)[:limit]

    # ----------------------------
    # Main collector
    # ----------------------------
    def _project_reference_publication_ids(self):
        if not self.project:
            return []
        return list(
            PublicationReference.objects.filter(
                publication__publicationproject__project=self.project,
                publication__private=False,
                referenced_publication__isnull=False,
                referenced_publication__private=False,
            )
            .values_list("referenced_publication_id", flat=True)
            .distinct()
        )

    def _project_publication_file_queryset(self):
        if not self.project:
            return PublicationFile.objects.none()
        reference_publication_ids = self._project_reference_publication_ids()
        return (
            PublicationFile.objects.filter(
                Q(
                    publication__publicationproject__project=self.project,
                    publication__private=False,
                )
                | Q(publication_id__in=reference_publication_ids)
            )
            .select_related("publication")
            .distinct()
        )

    @staticmethod
    def _normalize_target_type(raw_type):
        value = str(raw_type or "").strip().lower()
        if value in {"project", "project_document", "document", "project_file"}:
            return "project"
        if value in {"publication_file", "publicationfile", "pub_file", "publication"}:
            return "publication_file"
        if value in {"auto", "project/publication_file", "project_or_publication_file"}:
            return "auto"
        return ""

    def _project_document_by_id(self, document_id):
        if not self.project:
            return None
        try:
            parsed_id = int(document_id or 0)
        except (TypeError, ValueError):
            return None
        if parsed_id <= 0:
            return None
        return self.project.documents.filter(id=parsed_id).first()

    def _project_publication_file_by_id(self, file_id):
        try:
            parsed_id = int(file_id or 0)
        except (TypeError, ValueError):
            return None
        if parsed_id <= 0:
            return None
        return self._project_publication_file_queryset().filter(id=parsed_id).first()

    def _build_detailed_project_document_payload(self, document, text_limit):
        if not document:
            return None
        parsed_content = self.parse_document_content(document, text_limit=text_limit)
        return {
            "id": document.id,
            "type": "project",
            "title": self._clean_text(document.title),
            "file_type": self._clean_text(document.file_type).lower(),
            "version": document.version,
            "content": parsed_content,
            "created_at": document.created_at.isoformat() if document.created_at else None,
            "updated_at": document.updated_at.isoformat() if document.updated_at else None,
            "file_url": document.file.url if document.file else None,
        }

    def _build_detailed_publication_file_payload(self, publication_file, text_limit):
        if not publication_file:
            return None

        publication = publication_file.publication
        publication_title = self._clean_text(getattr(publication, "title_original", ""))
        file_name = ""
        if getattr(publication_file, "file", None):
            file_name = Path(getattr(publication_file.file, "name", "") or "").name
        title_parts = ["Файл публикации"]
        if publication_title:
            title_parts.append(publication_title)
        if file_name:
            title_parts.append(file_name)

        parsed_content = self.parse_document_content(publication_file, text_limit=text_limit)
        return {
            "id": publication_file.id,
            "type": "publication_file",
            "title": " — ".join(title_parts),
            "file_type": self._resolve_document_type(publication_file),
            "version": None,
            "content": parsed_content,
            "created_at": publication.created_at.isoformat() if publication and publication.created_at else None,
            "updated_at": publication.updated_at.isoformat() if publication and publication.updated_at else None,
            "file_url": publication_file.file.url if publication_file.file else None,
            "publication_id": publication.id if publication else None,
            "publication_title": publication_title,
            "publication_year": publication.year if publication else None,
        }

    def collect_detailed_documents(self, targets, active_document_id=None, active_publication_file_id=None, text_limit=120000):
        documents = []
        requested_targets = []
        seen_keys = set()

        def add_project_document(document, reason):
            payload = self._build_detailed_project_document_payload(document, text_limit=text_limit)
            if not payload:
                return False
            dedupe_key = f"project:{payload['id']}"
            if dedupe_key in seen_keys:
                return False
            seen_keys.add(dedupe_key)
            payload["reason"] = reason
            documents.append(payload)
            return True

        def add_publication_file(publication_file, reason):
            payload = self._build_detailed_publication_file_payload(publication_file, text_limit=text_limit)
            if not payload:
                return False
            dedupe_key = f"publication_file:{payload['id']}"
            if dedupe_key in seen_keys:
                return False
            seen_keys.add(dedupe_key)
            payload["reason"] = reason
            documents.append(payload)
            return True

        active_project_document = self._project_document_by_id(active_document_id)
        if active_project_document:
            add_project_document(active_project_document, reason="active_file")

        active_publication_file = self._project_publication_file_by_id(active_publication_file_id)
        if active_publication_file:
            add_publication_file(active_publication_file, reason="active_file")

        for raw_target in targets or []:
            if not isinstance(raw_target, dict):
                continue
            target_id = raw_target.get("id")
            target_type = self._normalize_target_type(raw_target.get("type"))
            try:
                parsed_id = int(target_id or 0)
            except (TypeError, ValueError):
                continue
            if parsed_id <= 0 or not target_type:
                continue

            requested_targets.append({"id": parsed_id, "type": target_type})
            if target_type == "project":
                add_project_document(self._project_document_by_id(parsed_id), reason="selected")
                continue
            if target_type == "publication_file":
                add_publication_file(self._project_publication_file_by_id(parsed_id), reason="selected")
                continue

            # "auto": try both sources in stable order.
            added = add_project_document(self._project_document_by_id(parsed_id), reason="selected")
            if not added:
                add_publication_file(self._project_publication_file_by_id(parsed_id), reason="selected")

        return {
            "requested_targets": requested_targets,
            "documents": documents,
        }

    def get_publication_file_context_by_id(self, publication_file_id):
        publication_file = self._project_publication_file_by_id(publication_file_id)
        if not publication_file:
            return ""
        last_document_text_limit = self.document_text_limit
        self.document_text_limit = 600000  # ~240k tokens reserve
        context = self.parse_document_content(publication_file) if publication_file else ""
        self.document_text_limit = last_document_text_limit
        return context

    def collect_project_documents(self):
        document_context = {}

        if not self.project:
            return document_context

        for document in self.project.documents.all():
            parsed_content = self.parse_document_content(document)

            document_context[document.id] = {
                "id": document.id,
                "title": document.title,
                "file_type": document.file_type,
                "version": document.version,
                "content": parsed_content,   # context limit
                "created_at": document.created_at.isoformat() if document.created_at else None,
                "updated_at": document.updated_at.isoformat() if document.updated_at else None,
                "file_url": document.file.url if document.file else None,
                "source_type": "project_document",
            }

        project_publication_documents = (
            self._project_publication_file_queryset()
            .filter(publication__publicationproject__project=self.project)
            .order_by("publication_id", "id")
        )

        reference_publication_ids = self._project_reference_publication_ids()
        reference_documents = self._project_publication_file_queryset().filter(
            publication_id__in=reference_publication_ids
        ).order_by("publication_id", "id")

        added_publication_file_ids = set()

        def _append_publication_file(publication_file, source_type):
            if publication_file.id in added_publication_file_ids:
                return
            added_publication_file_ids.add(publication_file.id)

            publication = publication_file.publication
            parsed_content = self.parse_document_content(publication_file)
            resolved_file_type = self._resolve_document_type(publication_file)
            file_name = ""
            if getattr(publication_file, "file", None):
                file_name = Path(getattr(publication_file.file, "name", "") or "").name

            title_prefix = "Файл публикации"
            if source_type == "reference_publication_file":
                title_prefix = "Файл reference-публикации"
            title_parts = [title_prefix]
            publication_title = self._clean_text(getattr(publication, "title_original", ""))
            if publication_title:
                title_parts.append(publication_title)
            if file_name:
                title_parts.append(file_name)

            document_context[f"publication_file:{publication_file.id}"] = {
                "id": publication_file.id,
                "title": " — ".join(title_parts),
                "file_type": resolved_file_type,
                "version": None,
                "content": parsed_content,
                "created_at": publication.created_at.isoformat() if publication and publication.created_at else None,
                "updated_at": publication.updated_at.isoformat() if publication and publication.updated_at else None,
                "file_url": publication_file.file.url if publication_file.file else None,
                "source_type": source_type,
                "publication_id": publication.id if publication else None,
                "publication_title": publication_title,
                "publication_year": publication.year if publication else None,
            }

        for publication_file in project_publication_documents:
            _append_publication_file(publication_file, source_type="project_publication_file")

        for publication_file in reference_documents:
            _append_publication_file(publication_file, source_type="reference_publication_file")

        return document_context

    def collect_reference(self):
        if not self.project:
            return {
                "total_references": 0,
                "publications": [],
            }

        reference_entries = (
            PublicationReference.objects.filter(
                publication__publicationproject__project=self.project,
                publication__private=False,
            )
            .select_related(
                "publication",
                "publication__venue",
                "referenced_publication",
                "referenced_publication__venue",
            )
            .order_by("publication__title_original", "publication_id", "order", "id")
        )

        grouped_by_publication = {}

        for reference in reference_entries:
            publication = reference.publication
            publication_group = grouped_by_publication.get(publication.id)
            if publication_group is None:
                publication_group = {
                    "publication_id": publication.id,
                    "title": self._clean_text(publication.title_original),
                    "year": publication.year,
                    "venue": self._clean_text(publication.venue.name) if publication.venue_id and publication.venue else "",
                    "references": [],
                }
                grouped_by_publication[publication.id] = publication_group

            referenced_publication = reference.referenced_publication
            linked_publication_payload = None
            if referenced_publication and not referenced_publication.private:
                linked_publication_payload = {
                    "id": referenced_publication.id,
                    "title": self._clean_text(referenced_publication.title_original),
                    "year": referenced_publication.year,
                    "venue": self._clean_text(referenced_publication.venue.name)
                    if referenced_publication.venue_id and referenced_publication.venue
                    else "",
                    "doi": self._clean_text(referenced_publication.doi),
                }

            raw_text = self._clean_text(reference.raw_text)
            display_text = raw_text
            if not display_text:
                if linked_publication_payload:
                    display_text = reference.display_text
                elif referenced_publication and referenced_publication.private:
                    display_text = "Private linked publication"
                else:
                    display_text = self._clean_text(reference.display_text)

            publication_group["references"].append(
                {
                    "id": reference.id,
                    "order": reference.order,
                    "note": self._clean_text(reference.note),
                    "raw_text": raw_text,
                    "display_text": self._clean_text(display_text),
                    "referenced_publication": linked_publication_payload,
                }
            )

        publications = list(grouped_by_publication.values())
        total_references = sum(len(item["references"]) for item in publications)
        return {
            "total_references": total_references,
            "publications": publications,
        }

    # ----------------------------
    # Build full context
    # ----------------------------
    def build(self):
        return {
            "project": {
                "id": self.project.id,
                "title": self.project.name,
            } if self.project else None,

            "user": {
                "id": self.user.id,
                "username": self.user.username,
            } if self.user else None,

            "documents": self.collect_project_documents(),
            "references": self.collect_reference(),
        }
    
    def get_document_context_by_id(self, document_id):
        document = self.project.documents.filter(id=document_id).first()
        last_document_text_limit = self.document_text_limit
        self.document_text_limit = 600000 # 240 000 токен (резерв для документа)
        context = self.parse_document_content(document) if document else ""
        self.document_text_limit = last_document_text_limit 
        return context
 
