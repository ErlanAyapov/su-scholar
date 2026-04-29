from __future__ import annotations

import tempfile
from pathlib import Path
from unittest import TestCase
from unittest.mock import patch

from docx import Document

from llm.agent.document_editor import DocumentEditor
from llm.agent.formula_utils import has_latex_formula_markers, split_text_and_formulas
from llm.agent.orchestrator import ProjectAgentOrchestrator


class FormulaUtilsTests(TestCase):
    def test_has_latex_formula_markers(self):
        self.assertFalse(has_latex_formula_markers("plain text"))
        self.assertTrue(has_latex_formula_markers("$x+y$"))
        self.assertTrue(has_latex_formula_markers("$$x+y$$"))

    def test_split_text_without_formulas(self):
        self.assertEqual(
            split_text_and_formulas("plain text"),
            [{"type": "text", "text": "plain text"}],
        )

    def test_split_text_with_one_inline_formula(self):
        self.assertEqual(
            split_text_and_formulas("value is $x+y$ now"),
            [
                {"type": "text", "text": "value is "},
                {"type": "inline_formula", "latex": "x+y"},
                {"type": "text", "text": " now"},
            ],
        )

    def test_split_text_with_inline_and_block_formulas(self):
        self.assertEqual(
            split_text_and_formulas("inline $x$ block $$y^2$$ end"),
            [
                {"type": "text", "text": "inline "},
                {"type": "inline_formula", "latex": "x"},
                {"type": "text", "text": " block "},
                {"type": "block_formula", "latex": "y^2"},
                {"type": "text", "text": " end"},
            ],
        )

    def test_split_text_with_several_formulas(self):
        segments = split_text_and_formulas("$a$ + $b$ = $$c$$")
        self.assertEqual([segment["type"] for segment in segments], [
            "inline_formula",
            "text",
            "inline_formula",
            "text",
            "block_formula",
        ])


class DocumentEditorFormulaTests(TestCase):
    def _create_docx(self, directory: str) -> Path:
        path = Path(directory) / "source.docx"
        document = Document()
        document.add_paragraph("First paragraph")
        document.add_paragraph("Second paragraph")
        document.save(path)
        return path

    def test_insert_after_plain_text_still_works(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = self._create_docx(temp_dir)
            editor = DocumentEditor(str(path))
            applied = editor.apply_operations([
                {"op": "insert_after", "paragraph_index": 0, "new_text": "Inserted plain text"}
            ])
            editor.save()

            self.assertEqual(applied, 1)
            self.assertEqual(Document(path).paragraphs[1].text, "Inserted plain text")

    def test_replace_plain_text_still_works(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = self._create_docx(temp_dir)
            editor = DocumentEditor(str(path))
            applied = editor.apply_operations([
                {
                    "op": "replace",
                    "paragraph_index": 0,
                    "old_text": "First paragraph",
                    "new_text": "Updated paragraph",
                }
            ])
            editor.save()

            self.assertEqual(applied, 1)
            self.assertEqual(Document(path).paragraphs[0].text, "Updated paragraph")

    def test_insert_after_formula_text_does_not_break_docx(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = self._create_docx(temp_dir)
            editor = DocumentEditor(str(path))
            applied = editor.apply_operations([
                {"op": "insert_after", "paragraph_index": 0, "new_text": "Formula $x+y$"}
            ])
            editor.save()

            self.assertEqual(applied, 1)
            self.assertGreaterEqual(len(Document(path).paragraphs), 3)

    def test_replace_formula_text_does_not_break_docx(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = self._create_docx(temp_dir)
            editor = DocumentEditor(str(path))
            applied = editor.apply_operations([
                {
                    "op": "replace",
                    "paragraph_index": 0,
                    "old_text": "First paragraph",
                    "new_text": "Formula $x+y$",
                }
            ])
            editor.save()

            self.assertEqual(applied, 1)
            self.assertGreaterEqual(len(Document(path).paragraphs), 2)

    def test_formula_conversion_failure_falls_back_to_plain_text(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = self._create_docx(temp_dir)
            editor = DocumentEditor(str(path))
            with patch(
                "llm.agent.formula_utils.latex_to_mathml_string",
                side_effect=RuntimeError("conversion failed"),
            ):
                applied = editor.apply_operations([
                    {"op": "insert_after", "paragraph_index": 0, "new_text": "Formula $x+y$"}
                ])
                editor.save()

            self.assertEqual(applied, 1)
            self.assertEqual(Document(path).paragraphs[1].text, "Formula $x+y$")

    def test_formula_replace_failure_preserves_surrounding_text(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = self._create_docx(temp_dir)
            editor = DocumentEditor(str(path))
            with patch(
                "llm.agent.formula_utils.latex_to_mathml_string",
                side_effect=RuntimeError("conversion failed"),
            ):
                applied = editor.apply_operations([
                    {
                        "op": "replace",
                        "paragraph_index": 0,
                        "old_text": "paragraph",
                        "new_text": "$x+y$",
                    }
                ])
                editor.save()

            self.assertEqual(applied, 1)
            self.assertEqual(Document(path).paragraphs[0].text, "First $x+y$")


class ProjectAgentOrchestratorFormulaRoutingTests(TestCase):
    def test_node_realtime_without_formulas_stays_live(self):
        engine = ProjectAgentOrchestrator._document_operation_engine_for_plan(
            "node_realtime",
            [{"op": "insert_after", "paragraph_index": 0, "new_text": "plain text"}],
        )
        self.assertEqual(engine, "node_realtime")

    def test_node_realtime_with_formula_switches_to_docx_patch(self):
        engine = ProjectAgentOrchestrator._document_operation_engine_for_plan(
            "node_realtime",
            [{"op": "insert_after", "paragraph_index": 0, "new_text": "Formula $x+y$"}],
        )
        self.assertEqual(engine, "docx_patch")

    def test_docx_patch_engine_stays_docx_patch(self):
        engine = ProjectAgentOrchestrator._document_operation_engine_for_plan(
            "docx_patch",
            [{"op": "insert_after", "paragraph_index": 0, "new_text": "Formula $x+y$"}],
        )
        self.assertEqual(engine, "docx_patch")
